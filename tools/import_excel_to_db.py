#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
ورود Excel یکدست‌شده به دیتابیس PostgreSQL (Supabase) ربات

اجرا از ریشه پروژه:
    python tools/import_excel_to_db.py namasha_mathsbu_full_clean.xlsx
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

try:
    import psycopg
except ImportError:
    print("psycopg نصب نیست: pip install 'psycopg[binary]'")
    sys.exit(1)

try:
    from dotenv import load_dotenv

    # .env در ریشه پروژه
    load_dotenv(Path(__file__).resolve().parent.parent / ".env")
    load_dotenv()
except ImportError:
    pass


def fix_db_url(url: str) -> str:
    if not url:
        return url
    if "sslmode=" not in url and ("supabase.com" in url or "pooler.supabase" in url):
        join = "&" if "?" in url else "?"
        url = f"{url}{join}sslmode=require"
    return url



PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")
ARABIC_LETTERS = str.maketrans({"ي": "ی", "ى": "ی", "ك": "ک", "ة": "ه"})


def normalize_text(text: str) -> str:
    if not text:
        return text or ""
    import re
    s = str(text).translate(PERSIAN_DIGITS).translate(ARABIC_LETTERS)
    s = re.sub(r"[\u064B-\u065F\u0670\u0640]", "", s)
    return s.strip()


def get_or_create_course(conn, name: str) -> int:
    name = normalize_text(name or "") or "بدون نام"
    with conn.cursor() as cur:
        cur.execute("SELECT id FROM courses WHERE name = %s", (name,))
        row = cur.fetchone()
        if row:
            return row[0]
        cur.execute(
            """
            INSERT INTO courses (name, teacher, created_at)
            VALUES (%s, %s, %s)
            RETURNING id
            """,
            (name, None, datetime.now().isoformat()),
        )
        return cur.fetchone()[0]


def video_exists(conn, namasha_url: str) -> bool:
    with conn.cursor() as cur:
        cur.execute(
            "SELECT 1 FROM videos WHERE namasha_url = %s LIMIT 1",
            (namasha_url,),
        )
        return cur.fetchone() is not None


def insert_video(
    conn,
    course_id: int,
    title: str,
    episode: str,
    namasha_url: str,
    download_url: str,
):
    with conn.cursor() as cur:
        cur.execute(
            """
            INSERT INTO videos
            (course_id, title, episode, namasha_url, download_url, telegram_file_id, created_at)
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (
                course_id,
                normalize_text(title) or "بدون عنوان",
                normalize_text(episode or ""),
                namasha_url,
                download_url or None,
                None,
                datetime.now().isoformat(),
            ),
        )


def main() -> int:
    parser = argparse.ArgumentParser(description="Import Excel into bot PostgreSQL")
    parser.add_argument("excel", type=Path, help="path to cleaned xlsx")
    parser.add_argument("--database-url", default=os.getenv("DATABASE_URL"))
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    db_url = fix_db_url(args.database_url or "")
    if not db_url:
        print("❌ DATABASE_URL تنظیم نشده.")
        print("   در فایل .env پروژه بگذار یا با --database-url بده.")
        return 1
    if not args.excel.exists():
        print(f"❌ فایل پیدا نشد: {args.excel}")
        return 1

    wb = load_workbook(args.excel, read_only=True, data_only=True)
    sheet = wb["videos"] if "videos" in wb.sheetnames else wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        print("فایل خالی است.")
        return 1

    header = [str(c or "").strip() for c in rows[0]]

    def col(*names):
        for n in names:
            if n in header:
                return header.index(n)
        return None

    i_course = col("course_name", "course")
    i_title = col("title")
    i_ep = col("episode")
    i_url = col("namasha_url", "url")
    i_dl = col("download_url", "download")
    if i_url is None:
        print("❌ ستون namasha_url در Excel نیست.")
        return 1

    data = []
    for raw in rows[1:]:
        if not raw or all(c is None or str(c).strip() == "" for c in raw):
            continue
        url = str(raw[i_url] or "").strip()
        if not url:
            continue
        data.append(
            {
                "course": str(raw[i_course] or "").strip() if i_course is not None else "",
                "title": str(raw[i_title] or "").strip() if i_title is not None else "",
                "episode": str(raw[i_ep] or "").strip() if i_ep is not None else "",
                "url": url,
                "download": str(raw[i_dl] or "").strip() if i_dl is not None else "",
            }
        )

    print(f"ردیف‌های معتبر Excel: {len(data)}")
    if args.dry_run:
        print("dry-run: چیزی نوشته نشد.")
        return 0

    print("اتصال به دیتابیس...")
    added = skipped = 0
    courses_seen = set()
    with psycopg.connect(db_url) as conn:
        for i, row in enumerate(data, 1):
            if video_exists(conn, row["url"]):
                skipped += 1
                continue
            course_id = get_or_create_course(conn, row["course"])
            courses_seen.add(row["course"])
            insert_video(
                conn,
                course_id,
                row["title"] or row["url"],
                row["episode"],
                row["url"],
                row["download"],
            )
            added += 1
            if i % 50 == 0:
                conn.commit()
                print(f"  ... {i}/{len(data)} (اضافه: {added}, رد تکراری: {skipped})")
        conn.commit()

    print(f"\n✅ تمام.")
    print(f"   ویدیو اضافه شد: {added}")
    print(f"   تکراری رد شد: {skipped}")
    print(f"   درس درگیر در این ورود: حدود {len(courses_seen)}")
    print("   فایل تلگرام (file_id) خالی است؛ بعداً با فوروارد وصل کن.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
