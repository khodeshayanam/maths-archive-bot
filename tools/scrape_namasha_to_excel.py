#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
جمع‌آوری ویدیوهای کانال mathsbu از نماشا و ذخیره در Excel

نصب (یک‌بار):
    pip install httpx beautifulsoup4 openpyxl lxml

اجرا:
    python tools/scrape_namasha_to_excel.py

خروجی:
    namasha_mathsbu.xlsx  در همان پوشه‌ای که دستور را زدی
"""

from __future__ import annotations

import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin, urlparse

import httpx
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

CHANNEL_URL = "https://www.namasha.com/mathsbu"
BASE = "https://www.namasha.com"
OUTPUT = Path("namasha_mathsbu.xlsx")

# فاصله بین درخواست‌ها (ثانیه) تا به سایت فشار نیاید
DELAY = 0.8
TIMEOUT = 40.0
MAX_PAGES = 200  # سقف ایمنی برای صفحه‌بندی
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.8",
}

# الگوهای عنوان فارسی: «نام درس - قسمت ۱۲» یا «نام درس قسمت 12»
EP_SPLIT = re.compile(
    r"\s*[-–—]?\s*قسمت\s*([0-9۰-۹]+(?:\s*[-–]\s*[0-9۰-۹]+)?)\s*$",
    re.IGNORECASE,
)
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")


def normalize_digits(s: str) -> str:
    return (s or "").translate(PERSIAN_DIGITS)


def parse_title(full_title: str) -> tuple[str, str, str]:
    """برمی‌گرداند (نام_درس، عنوان_کامل، شماره_قسمت)."""
    title = (full_title or "").strip()
    title = re.sub(r"\s+", " ", title)
    m = EP_SPLIT.search(title)
    if m:
        course = title[: m.start()].strip(" -–—")
        episode = normalize_digits(m.group(1).replace(" ", ""))
        return course or title, title, episode
    # fallback: آخرین عدد در عنوان
    # آخرین عدد را فقط اگر شماره قسمت معقول باشد (نه سال مثل ۱۴۰۲)
    nums = re.findall(r"[0-9۰-۹]+", title)
    episode = ""
    if nums:
        cand = normalize_digits(nums[-1])
        try:
            n = int(cand.split("-")[0])
            if 1 <= n <= 80:
                episode = cand
        except ValueError:
            episode = ""
    return title, title, episode


def abs_url(href: str) -> str:
    if not href:
        return ""
    return urljoin(BASE, href)


def extract_video_links_from_list_html(html: str) -> list[tuple[str, str]]:
    """لیست (url, title) از صفحه کانال."""
    soup = BeautifulSoup(html, "lxml")
    found: list[tuple[str, str]] = []
    seen = set()

    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not re.search(r"/v/[A-Za-z0-9_-]+", href):
            continue
        url = abs_url(href).split("?")[0].rstrip("/")
        if url in seen:
            continue
        # عنوان از متن لینک یا heading نزدیک
        text = a.get_text(" ", strip=True)
        if not text or len(text) < 3:
            parent = a.find_parent(["h1", "h2", "h3", "h4", "div", "li", "article"])
            if parent:
                text = parent.get_text(" ", strip=True)
        text = re.sub(r"\s+", " ", text or "")
        # حذف مدت زمان از اول عنوان اگر چسبیده باشد
        text = re.sub(r"^[\d۰-۹:：]+\s*", "", text)
        if len(text) < 3:
            continue
        seen.add(url)
        found.append((url, text))
    return found


def find_next_page_url(html: str, current_url: str) -> str | None:
    soup = BeautifulSoup(html, "lxml")
    # rel=next
    link = soup.find("a", attrs={"rel": "next"})
    if link and link.get("href"):
        return abs_url(link["href"])
    # متن «بیشتر» / load more گاهی دکمه است نه لینک
    for a in soup.find_all("a", href=True):
        t = a.get_text(" ", strip=True)
        if t in ("بعدی", "بیشتر", "ویدیوهای بیشتر", "بارگذاری ویدیوهای بیشتر", "Next"):
            return abs_url(a["href"])
    # ?page=N
    m = re.search(r"[?&]page=(\d+)", current_url)
    if m:
        n = int(m.group(1)) + 1
        return re.sub(r"([?&])page=\d+", rf"\1page={n}", current_url)
    # اگر صفحه اول بود
    if "page=" not in current_url:
        join = "&" if "?" in current_url else "?"
        return f"{current_url}{join}page=2"
    return None


def extract_download_url(html: str, page_url: str) -> str:
    """تلاش برای پیدا کردن لینک مستقیم دانلود از HTML صفحه ویدیو."""
    # لینک‌های شناخته‌شده نماشا
    patterns = [
        r'https?://(?:www\.)?namasha\.com/videos/dl/[^"\'\s<>]+',
        r'https?://[^"\'\s<>]*namasha[^"\'\s<>]*/(?:dl|download)/[^"\'\s<>]+',
        r'"(https?://[^"]+\.mp4[^"]*)"',
        r"'(https?://[^']+\.mp4[^']*)'",
    ]
    for pat in patterns:
        m = re.search(pat, html, re.I)
        if m:
            url = m.group(1) if m.lastindex else m.group(0)
            url = url.replace("&amp;", "&")
            if "namasha" in url or url.endswith(".mp4"):
                return url
    soup = BeautifulSoup(html, "lxml")
    for a in soup.find_all("a", href=True):
        href = a["href"]
        label = a.get_text(" ", strip=True)
        if "/videos/dl/" in href or "download" in href.lower() or "دانلود" in label:
            return abs_url(href)
    return ""


def scrape_video_page(client: httpx.Client, url: str) -> dict:
    r = client.get(url, follow_redirects=True)
    r.raise_for_status()
    html = r.text
    soup = BeautifulSoup(html, "lxml")

    title = ""
    if soup.title and soup.title.string:
        title = soup.title.string.strip()
        title = re.sub(r"\s*[\|\-–]\s*نماشا\s*$", "", title).strip()
    og = soup.find("meta", property="og:title")
    if og and og.get("content"):
        title = og["content"].strip()
    h1 = soup.find(["h1", "h2"])
    if h1:
        t = h1.get_text(" ", strip=True)
        if t and len(t) > 3:
            title = t

    course, full_title, episode = parse_title(title)
    download = extract_download_url(html, url)

    return {
        "course_name": course,
        "title": full_title or title,
        "episode": episode,
        "namasha_url": url,
        "download_url": download,
    }


def collect_all_list_urls(client: httpx.Client) -> list[tuple[str, str]]:
    """پیمایش صفحات کانال و جمع‌آوری لینک ویدیوها."""
    all_items: list[tuple[str, str]] = []
    seen_urls = set()
    url = CHANNEL_URL
    empty_streak = 0

    for page_i in range(1, MAX_PAGES + 1):
        print(f"  صفحه لیست {page_i}: {url}")
        try:
            r = client.get(url, follow_redirects=True)
            r.raise_for_status()
        except Exception as e:
            print(f"  ⚠️ خطا در دریافت لیست: {e}")
            break

        batch = extract_video_links_from_list_html(r.text)
        new = 0
        for vurl, title in batch:
            if vurl not in seen_urls:
                seen_urls.add(vurl)
                all_items.append((vurl, title))
                new += 1

        print(f"    → {new} ویدیوی جدید (جمع: {len(all_items)})")
        if new == 0:
            empty_streak += 1
            if empty_streak >= 2:
                break
        else:
            empty_streak = 0

        next_url = find_next_page_url(r.text, url)
        if not next_url or next_url == url:
            # تلاش با page= صریح
            if page_i == 1 and "page=" not in url:
                next_url = f"{CHANNEL_URL}?page=2"
            else:
                break
        url = next_url
        time.sleep(DELAY)

    return all_items


def save_excel(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "videos"
    headers = [
        "course_name",
        "title",
        "episode",
        "namasha_url",
        "download_url",
    ]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")

    for row in rows:
        ws.append([row.get(h, "") for h in headers])

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 50
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 45
    ws.column_dimensions["E"].width = 55

    # برگه راهنما
    info = wb.create_sheet("README", 0)
    info["A1"] = "راهنمای فایل"
    info["A1"].font = Font(bold=True, size=14)
    lines = [
        "",
        "ستون course_name: نام درس (برای گروه‌بندی در ربات)",
        "ستون title: عنوان کامل ویدیو",
        "ستون episode: شماره قسمت (برای مرتب‌سازی)",
        "ستون namasha_url: لینک تماشا",
        "ستون download_url: لینک دانلود مستقیم (اگر پیدا شده؛ وگرنه خالی)",
        "",
        "می‌توانید قبل از ورود به دیتابیس، نام درس‌ها را در Excel یکدست کنید.",
        "مرحله بعد: اسکریپت import_excel_to_db.py",
    ]
    for i, line in enumerate(lines, start=2):
        info[f"A{i}"] = line
    info.column_dimensions["A"].width = 80

    wb.save(path)


def main() -> int:
    print("شروع جمع‌آوری از:", CHANNEL_URL)
    print("اگر فیلتر است، VPN را روشن کنید.\n")

    rows: list[dict] = []
    with httpx.Client(headers=HEADERS, timeout=TIMEOUT) as client:
        print("۱) خواندن لیست ویدیوهای کانال...")
        items = collect_all_list_urls(client)
        if not items:
            print("❌ هیچ ویدیویی پیدا نشد. VPN / ساختار سایت را بررسی کنید.")
            return 1

        print(f"\n۲) دریافت جزئیات {len(items)} ویدیو...")
        for i, (url, list_title) in enumerate(items, 1):
            try:
                data = scrape_video_page(client, url)
                # اگر عنوان صفحه ضعیف بود از لیست استفاده کن
                if not data["title"] or len(data["title"]) < 4:
                    course, full, ep = parse_title(list_title)
                    data.update(
                        course_name=course,
                        title=full,
                        episode=ep,
                        namasha_url=url,
                    )
                elif not data["course_name"]:
                    course, _, ep = parse_title(data["title"] or list_title)
                    data["course_name"] = course
                    if not data["episode"]:
                        data["episode"] = ep
                rows.append(data)
                flag = "↓" if data.get("download_url") else "·"
                print(f"  [{i}/{len(items)}] {flag} {data.get('title', url)[:70]}")
            except Exception as e:
                print(f"  [{i}/{len(items)}] ⚠️ {url}: {e}")
                course, full, ep = parse_title(list_title)
                rows.append(
                    {
                        "course_name": course,
                        "title": full or list_title,
                        "episode": ep,
                        "namasha_url": url,
                        "download_url": "",
                    }
                )
            time.sleep(DELAY)

    save_excel(rows, OUTPUT)
    with_dl = sum(1 for r in rows if r.get("download_url"))
    print(f"\n✅ ذخیره شد: {OUTPUT.resolve()}")
    print(f"   تعداد ردیف: {len(rows)}")
    print(f"   دارای لینک دانلود: {with_dl}")
    print("\nقدم بعد: یکدست‌کردن نام درس‌ها در Excel، سپس import به دیتابیس.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
