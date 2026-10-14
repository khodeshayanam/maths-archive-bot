#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
جمع‌آوری کامل آرشیو کانال mathsbu از نماشا → Excel
همه درس‌ها از جمله داده کاوی دکتر دهقان.

نصب یک‌بار:
    pip install httpx beautifulsoup4 lxml openpyxl playwright

اجرا (با Chrome ویندوز؛ VPN اگر لازم است):
    python tools/scrape_full_archive.py

خروجی:
    namasha_mathsbu_complete.xlsx
"""

from __future__ import annotations

import argparse
import re
import sys
import time
from pathlib import Path
from urllib.parse import urljoin

import httpx
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

CHANNEL_URL = "https://www.namasha.com/mathsbu"
BASE = "https://www.namasha.com"
OUTPUT = Path("namasha_mathsbu_complete.xlsx")
DELAY = 0.5
TIMEOUT = 45.0

# پیش‌فرض: هیچ درسی حذف نمی‌شود (همه آرشیو)
DEFAULT_EXCLUDE: list[str] = []

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/122.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "fa-IR,fa;q=0.9,en;q=0.8",
}

EP_SPLIT = re.compile(
    r"\s*[-–—]?\s*قسمت\s*([0-9۰-۹]+(?:\s*[-–]\s*[0-9۰-۹]+)?)\s*$",
    re.IGNORECASE,
)
PERSIAN_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹", "0123456789")


def normalize_digits(s: str) -> str:
    return (s or "").translate(PERSIAN_DIGITS)


def parse_title(full_title: str) -> tuple[str, str, str]:
    title = re.sub(r"\s+", " ", (full_title or "").strip())
    m = EP_SPLIT.search(title)
    if m:
        course = title[: m.start()].strip(" -–—")
        episode = normalize_digits(m.group(1).replace(" ", ""))
        return course or title, title, episode
    # آخرین عدد را فقط اگر شماره قسمت معقول باشد (نه سال مثل ۱۴۰۲)
    nums = re.findall(r"[0-9۰-۹]+", title)
    episode = ""
    if nums:
        cand = normalize_digits(nums[-1])
        try:
            n = int(cand.split("-")[0])
            if 1 <= n <= 80:
                episode = cand
        except ValueError:
            episode = ""
    return title, title, episode


def abs_url(href: str) -> str:
    return urljoin(BASE, href) if href else ""


def should_exclude(course: str, title: str, exclude_terms: list[str]) -> bool:
    blob = f"{course} {title}".replace("‌", " ")
    blob_norm = re.sub(r"\s+", " ", blob).strip().lower()
    for term in exclude_terms:
        t = re.sub(r"\s+", " ", term.replace("‌", " ")).strip().lower()
        if t and t in blob_norm:
            return True
    return False


def extract_video_links(html: str) -> list[tuple[str, str]]:
    soup = BeautifulSoup(html, "lxml")
    found: list[tuple[str, str]] = []
    seen: set[str] = set()
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not re.search(r"/v/[A-Za-z0-9_-]+", href):
            continue
        url = abs_url(href).split("?")[0].rstrip("/")
        if url in seen:
            continue
        text = a.get_text(" ", strip=True)
        if not text or len(text) < 3:
            parent = a.find_parent(["h1", "h2", "h3", "h4", "div", "li", "article"])
            if parent:
                text = parent.get_text(" ", strip=True)
        text = re.sub(r"\s+", " ", text or "")
        text = re.sub(r"^[\d۰-۹:：]+\s*", "", text)
        if len(text) < 3:
            continue
        seen.add(url)
        found.append((url, text))
    return found


def extract_download_url(html: str) -> str:
    patterns = [
        r'https?://(?:www\.)?namasha\.com/videos/dl/[^"\'\s<>]+',
        r'"(https?://[^"]+\.mp4[^"]*)"',
        r"'(https?://[^']+\.mp4[^']*)'",
    ]
    for pat in patterns:
        m = re.search(pat, html, re.I)
        if m:
            url = m.group(1) if m.lastindex else m.group(0)
            return url.replace("&amp;", "&")
    soup = BeautifulSoup(html, "lxml")
    for a in soup.find_all("a", href=True):
        href = a["href"]
        label = a.get_text(" ", strip=True)
        if "/videos/dl/" in href or "دانلود" in label:
            return abs_url(href)
    return ""


def scrape_video_page(client: httpx.Client, url: str, list_title: str) -> dict:
    r = client.get(url, follow_redirects=True)
    r.raise_for_status()
    html = r.text
    soup = BeautifulSoup(html, "lxml")
    title = list_title
    og = soup.find("meta", property="og:title")
    if og and og.get("content"):
        title = og["content"].strip()
    elif soup.title and soup.title.string:
        title = re.sub(r"\s*[\|\-–]\s*نماشا\s*$", "", soup.title.string.strip())
    h1 = soup.find(["h1", "h2"])
    if h1:
        t = h1.get_text(" ", strip=True)
        if t and len(t) > 3:
            title = t
    course, full_title, episode = parse_title(title)
    if not course or len(course) < 2:
        course, full_title, episode = parse_title(list_title)
    return {
        "course_name": course,
        "title": full_title or title or list_title,
        "episode": episode,
        "namasha_url": url,
        "download_url": extract_download_url(html),
    }


def _launch_browser(p):
    """
    اول Chrome/Edge خود ویندوز را امتحان می‌کند (بدون دانلود Chromium).
    اگر نبود، سراغ Chromiumِ Playwright می‌رود.
    """
    errors = []
    for channel in ("chrome", "msedge"):
        try:
            browser = p.chromium.launch(channel=channel, headless=True)
            print(f"  مرورگر: {channel}")
            return browser
        except Exception as e:
            errors.append(f"{channel}: {e}")
    try:
        browser = p.chromium.launch(headless=True)
        print("  مرورگر: chromium (playwright)")
        return browser
    except Exception as e:
        errors.append(f"chromium: {e}")
        msg = (
            "❌ هیچ مرورگری اجرا نشد.\n"
            "  • Google Chrome یا Microsoft Edge را نصب کن، یا\n"
            "  • با VPN پایدار بزن: playwright install chromium\n\n"
            + "\n".join(errors)
        )
        raise RuntimeError(msg) from e


def collect_links_with_browser(max_clicks: int = 80) -> list[tuple[str, str]]:
    """با Playwright دکمه «بارگذاری ویدیوهای بیشتر» را می‌زند تا لیست کامل شود."""
    try:
        from playwright.sync_api import sync_playwright
    except ImportError:
        print("❌ playwright نصب نیست.")
        print("   pip install playwright")
        sys.exit(1)

    print("باز کردن مرورگر و بارگذاری کل لیست کانال...")
    items: dict[str, str] = {}

    with sync_playwright() as p:
        browser = _launch_browser(p)
        page = browser.new_page(user_agent=HEADERS["User-Agent"])
        page.goto(CHANNEL_URL, wait_until="domcontentloaded", timeout=90000)
        page.wait_for_timeout(2500)

        for i in range(max_clicks):
            html = page.content()
            batch = extract_video_links(html)
            before = len(items)
            for url, title in batch:
                items[url] = title
            added = len(items) - before
            print(f"  کلیک {i + 1}: جمع لینک‌ها = {len(items)} (+{added})")

            # دکمه بارگذاری بیشتر
            clicked = False
            for label in (
                "بارگذاری ویدیوهای بیشتر",
                "ویدیوهای بیشتر",
                "بیشتر",
                "Load more",
            ):
                btn = page.get_by_text(label, exact=False)
                try:
                    if btn.count() > 0 and btn.first.is_visible():
                        btn.first.click(timeout=5000)
                        clicked = True
                        page.wait_for_timeout(2000)
                        break
                except Exception:
                    continue

            if not clicked:
                # اسکرول پایین به‌عنوان پشتیبان
                page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                page.wait_for_timeout(1500)
                html2 = page.content()
                batch2 = extract_video_links(html2)
                before2 = len(items)
                for url, title in batch2:
                    items[url] = title
                if len(items) == before2 and added == 0:
                    print("  دیگر ویدیوی جدیدی پیدا نشد — توقف.")
                    break

        browser.close()

    return list(items.items())


def collect_links_http_fallback() -> list[tuple[str, str]]:
    """فقط صفحه اول (+ تلاش page=) — ممکن است ناقص باشد."""
    print("حالت ساده (بدون playwright) — ممکن است فقط بخشی از آرشیو بیاید.")
    items: dict[str, str] = {}
    with httpx.Client(headers=HEADERS, timeout=TIMEOUT) as client:
        for page_i in range(1, 30):
            url = CHANNEL_URL if page_i == 1 else f"{CHANNEL_URL}?page={page_i}"
            try:
                r = client.get(url, follow_redirects=True)
                r.raise_for_status()
            except Exception as e:
                print(f"  خطا صفحه {page_i}: {e}")
                break
            batch = extract_video_links(r.text)
            before = len(items)
            for u, t in batch:
                items[u] = t
            print(f"  صفحه {page_i}: جمع = {len(items)} (+{len(items) - before})")
            if page_i > 1 and len(items) == before:
                break
            time.sleep(DELAY)
    return list(items.items())


def save_excel(rows: list[dict], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "videos"
    headers = ["course_name", "title", "episode", "namasha_url", "download_url"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 44
    ws.column_dimensions["E"].width = 56

    info = wb.create_sheet("README", 0)
    info["A1"] = "آرشیو کامل mathsbu (همه درس‌ها)"
    info["A1"].font = Font(bold=True, size=13)
    info["A3"] = "این فایل فقط Excel است — هنوز به دیتابیس وارد نشده."
    info["A4"] = f"تعداد ردیف: {len(rows)}"
    info["A5"] = "قدم بعد: python tools/import_excel_to_db.py namasha_mathsbu_complete.xlsx"
    info.column_dimensions["A"].width = 80
    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--exclude",
        action="append",
        default=None,
        help="عبارت حذف از نام درس/عنوان (قابل تکرار)",
    )
    parser.add_argument(
        "--no-browser",
        action="store_true",
        help="بدون Playwright (ممکن است ناقص باشد)",
    )
    parser.add_argument("--max-clicks", type=int, default=120)
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=OUTPUT,
        help="مسیر فایل Excel خروجی",
    )
    args = parser.parse_args()
    exclude = args.exclude if args.exclude is not None else list(DEFAULT_EXCLUDE)

    print("کانال:", CHANNEL_URL)
    print("حذف درس‌ها:", exclude if exclude else "(هیچ — همه درس‌ها)")
    print("خروجی:", args.output)
    print()

    if args.no_browser:
        link_items = collect_links_http_fallback()
    else:
        try:
            link_items = collect_links_with_browser(max_clicks=args.max_clicks)
        except Exception as e:
            print(f"⚠️ مرورگر شکست خورد ({e}) — تلاش حالت ساده...")
            link_items = collect_links_http_fallback()

    if not link_items:
        print("❌ هیچ لینکی پیدا نشد. VPN را چک کن.")
        return 1

    print(f"\nتعداد لینک یکتا: {len(link_items)}")
    print("دریافت جزئیات هر ویدیو...\n")

    rows: list[dict] = []
    skipped = 0
    with httpx.Client(headers=HEADERS, timeout=TIMEOUT) as client:
        for i, (url, list_title) in enumerate(link_items, 1):
            try:
                data = scrape_video_page(client, url, list_title)
            except Exception as e:
                print(f"  [{i}/{len(link_items)}] ⚠️ {url}: {e}")
                course, full, ep = parse_title(list_title)
                data = {
                    "course_name": course,
                    "title": full or list_title,
                    "episode": ep,
                    "namasha_url": url,
                    "download_url": "",
                }

            if should_exclude(data["course_name"], data["title"], exclude):
                skipped += 1
                print(f"  [{i}/{len(link_items)}] ⏭ حذف: {data['title'][:60]}")
            else:
                rows.append(data)
                mark = "↓" if data.get("download_url") else "·"
                print(f"  [{i}/{len(link_items)}] {mark} {data['title'][:70]}")
            time.sleep(DELAY)

    out = args.output
    save_excel(rows, out)
    print(f"\n✅ ذخیره شد: {out.resolve()}")
    print(f"   ردیف در Excel: {len(rows)}")
    print(f"   حذف‌شده با --exclude: {skipped}")
    print("   هنوز به دیتابیس چیزی وارد نشده.")
    print("   قدم بعد: python tools/import_excel_to_db.py", out.name)
    return 0


if __name__ == "__main__":
    sys.exit(main())
